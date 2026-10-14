"""
Excel multi-sheet exporter - DRASTIC-Tox v61
University of Khartoum, Faculty of Engineering

Sheets: Summary | DRASTIC Breakdown | SPSA | Monte Carlo | Metadata
Requires: openpyxl

Changes vs v58.9
  * ranges corrected: DRASTIC is 23-226 (not 230); DRASTIC-Tox is EXPERIMENTAL, range 23-276 with equal weights
  * two-axis result added (vulnerability class x hazard class), computed with drastic_core
  * limits/standard actually used (EU Directive 98/83/EC) are written to the report; WHO cyanide GV was withdrawn
  * numbers are written as numbers (were all text); text starting with = + - @ is neutralised (formula injection)
  * unit warnings and "assumed uncertainty" / "screening only" notes included
  * metadata references now match what the code really uses
"""
import datetime
from io import BytesIO

COLOR_PRIMARY, COLOR_ACCENT, COLOR_LIGHT_BG = "5C2C16", "C19A6B", "FAF8F3"
COLOR_LOW, COLOR_MEDIUM, COLOR_HIGH, COLOR_VERY_HIGH, COLOR_WHITE = "388E3C", "FBC02D", "F57C00", "D32F2F", "FFFFFF"

try:
    import drastic_core as dc
    from drastic_bridge import STANDARD, DRASTIC_MAX
except Exception:            # keep the exporter usable even if the core is missing
    dc, STANDARD, DRASTIC_MAX = None, "EU", 226


def _has_openpyxl():
    try:
        from openpyxl import Workbook
        return True, Workbook
    except ImportError:
        return False, None


def _safe(v):
    """Keep numbers numeric; neutralise spreadsheet formulas in text."""
    if isinstance(v, bool) or v is None:
        return "" if v is None else str(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def _num_or_text(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() and not isinstance(v, str) else f
    except (TypeError, ValueError):
        return _safe(v)


def _style_header(cell, bg_color=COLOR_PRIMARY, font_color=COLOR_WHITE, size=12):
    from openpyxl.styles import Font, PatternFill, Alignment
    cell.font = Font(bold=True, color=font_color, size=size)
    cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_title(cell, size=14):
    from openpyxl.styles import Font, Alignment
    cell.font = Font(bold=True, color=COLOR_PRIMARY, size=size)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _style_label(cell):
    from openpyxl.styles import Font, PatternFill, Alignment
    cell.font = Font(bold=True, color=COLOR_PRIMARY, size=10)
    cell.fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _style_value(cell, bold=False):
    from openpyxl.styles import Font, Alignment
    cell.font = Font(bold=bold, size=10)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _level_color(level_str):
    level = (level_str or "").lower()
    if "very_high" in level or "very high" in level or "مرتفع جد" in level or "critical" in level:
        return COLOR_VERY_HIGH
    if "high" in level or "مرتفع" in level:
        return COLOR_HIGH
    if "medium" in level or "moderate" in level or "متوسط" in level:
        return COLOR_MEDIUM
    return COLOR_LOW


def _kv(ws, row, label, val, bold=False):
    ws.cell(row=row, column=1, value=_safe(label))
    _style_label(ws.cell(row=row, column=1))
    ws.cell(row=row, column=2, value=_num_or_text(val))
    _style_value(ws.cell(row=row, column=2), bold=bold)
    return row + 1


def _section(ws, row, text, cols=2):
    ws.cell(row=row, column=1, value=text)
    _style_header(ws.cell(row=row, column=1))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    return row + 1


def _two_axis(site_info, drastic_result):
    """(vulnerability, hazard, risk) from drastic_core, or None if not computable."""
    if dc is None:
        return None
    try:
        base = float(drastic_result.get("base_drastic"))
        cn, hg = float(site_info.get("cn", 0)), float(site_info.get("hg", 0))
        L = dc.LIMITS_MG_L[STANDARD]
        hq = max(cn / L["cn"] if L["cn"] else 0.0, hg / L["hg"])
        v, vp = dc.vulnerability_class(base)
        h = dc.hazard_class(hq)
        return v, vp, h, round(hq, 1), dc.risk_category(v, h), L
    except Exception:
        return None


def generate_excel_report(site_info, drastic_result, risk_info, mining_type, spsa_result=None, mc_result=None):
    """Generate the Excel report. Returns bytes (None if openpyxl is missing)."""
    ok, Workbook = _has_openpyxl()
    if not ok:
        return None
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    now = datetime.datetime.now()

    # ------------------------------------------------------------ 1. Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "DRASTIC-Tox Report - Site Summary (screening level)"
    _style_title(ws["A1"], size=16)
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 30
    ws["A2"] = f"Generated: {now.strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(italic=True, size=9, color="666666")
    ws.merge_cells("A2:D2")

    row = _section(ws, 4, "SITE INFORMATION")
    for label, key in (("Site Name", "name"), ("State", "state"), ("Coordinates", "coords"), ("Depth (m)", "depth"),
                       ("Recharge (mm/yr)", "recharge"), ("Slope (%)", "slope"), ("Conductivity (m/day)", "conductivity"),
                       ("Aquifer", "aquifer"), ("Soil", "soil"), ("Vadose", "vadose"),
                       ("CN in water (mg/L)", "cn"), ("Hg in water (mg/L)", "hg")):
        row = _kv(ws, row, label, site_info.get(key, "N/A"))
    row = _kv(ws, row, "Mining Type", mining_type)

    row = _section(ws, row + 1, f"VULNERABILITY (DRASTIC, range 23-{DRASTIC_MAX})")
    row = _kv(ws, row, "DRASTIC Index", drastic_result.get("base_drastic", "N/A"), bold=True)

    ta = _two_axis(site_info, drastic_result)
    if ta:
        v, vp, h, hq, rk, L = ta
        row = _section(ws, row + 1, "TWO-AXIS ASSESSMENT (recommended)")
        row = _kv(ws, row, "Vulnerability class", f"{v} ({vp}% of range)")
        row = _kv(ws, row, "Hazard class (max conc/limit)", f"{h} (HQ = {hq})")
        row = _kv(ws, row, "Combined risk", rk, bold=True)
        row = _kv(ws, row, "Limits used", f"{STANDARD}: CN {L['cn']} mg/L, Hg {L['hg']} mg/L (EU Directive 98/83/EC)")

    row = _section(ws, row + 1, "DRASTIC-Tox (EXPERIMENTAL, range 23-276 with equal weights)")
    for label, key in (("CN Score", "cn_score"), ("Hg Score", "hg_score"), ("alpha (CN weight)", "alpha"),
                       ("beta (Hg weight)", "beta"), ("Source Factor (SF)", "source_factor"),
                       ("Final Bonus", "toxicity_bonus")):
        row = _kv(ws, row, label, drastic_result.get(key, "N/A"))
    row = _kv(ws, row, "DRASTIC-Tox Index", drastic_result.get("drastic_t", "N/A"), bold=True)
    row = _kv(ws, row, "Increase over DRASTIC", f"+{drastic_result.get('increase_pct', 0)}%")

    row = _section(ws, row + 1, "RISK LEVEL (index bands)")
    level_str = str(risk_info.get("level", "N/A"))
    ws.cell(row=row, column=1, value="Risk Level")
    _style_label(ws.cell(row=row, column=1))
    cell = ws.cell(row=row, column=2, value=_safe(level_str))
    cell.font = Font(bold=True, color=COLOR_WHITE, size=12)
    cell.fill = PatternFill(start_color=_level_color(level_str), end_color=_level_color(level_str), fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    row = _kv(ws, row + 1, "Recommended Action", risk_info.get("action", "N/A"))

    warnings = list(drastic_result.get("unit_warnings", []) or [])
    if warnings:
        row = _section(ws, row + 1, "DATA WARNINGS")
        for w in warnings:
            ws.cell(row=row, column=1, value=_safe(w)); ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            row += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15

    # ------------------------------------------------------------ 2. DRASTIC breakdown
    ws2 = wb.create_sheet("DRASTIC Breakdown")
    ws2["A1"] = "DRASTIC Parameter Breakdown (Aller et al. 1987 weights)"
    _style_title(ws2["A1"], size=16)
    ws2.merge_cells("A1:E1")
    ws2.row_dimensions[1].height = 30
    for col, h in enumerate(["Parameter", "Weight (Wi)", "Rating (Ri)", "Wi x Ri", "Contribution %"], 1):
        _style_header(ws2.cell(row=3, column=col, value=h), bg_color=COLOR_ACCENT, font_color=COLOR_PRIMARY, size=11)
    W = dc.W_STD if dc else dict(D=5, R=4, A=3, S=2, T=1, I=5, C=3)
    names = (("D", "D - Depth to Water"), ("R", "R - Net Recharge"), ("A", "A - Aquifer Media"), ("S", "S - Soil Media"),
             ("T", "T - Topography"), ("I", "I - Impact of Vadose"), ("C", "C - Hydraulic Conductivity"))
    rating = {k: site_info.get(f"{k}_rating", None) for k, _ in names}
    total = sum(W[k] * r for k, r in rating.items() if isinstance(r, (int, float)))
    row = 4
    for k, name in names:
        r = rating[k]
        ws2.cell(row=row, column=1, value=name)
        ws2.cell(row=row, column=2, value=W[k])
        ws2.cell(row=row, column=3, value=r if r is not None else "-")
        if isinstance(r, (int, float)):
            ws2.cell(row=row, column=4, value=W[k] * r)
            ws2.cell(row=row, column=5, value=round(W[k] * r / total * 100, 2) if total else 0)
        else:
            ws2.cell(row=row, column=4, value="-"); ws2.cell(row=row, column=5, value="-")
        for c in range(1, 6):
            _style_value(ws2.cell(row=row, column=c))
        row += 1
    row += 1
    for c, v in enumerate(["TOTAL DRASTIC", sum(W.values()), "-", total, 100 if total else "-"], 1):
        ws2.cell(row=row, column=c, value=v)
        _style_header(ws2.cell(row=row, column=c), bg_color=COLOR_ACCENT if c > 1 else COLOR_PRIMARY,
                      font_color=COLOR_PRIMARY if c > 1 else COLOR_WHITE)
    for col, w in (("A", 30), ("B", 15), ("C", 15), ("D", 15), ("E", 18)):
        ws2.column_dimensions[col].width = w

    # ------------------------------------------------------------ 3. SPSA
    if spsa_result and "parameters" in spsa_result:
        ws3 = wb.create_sheet("SPSA")
        ws3["A1"] = "Single Parameter Sensitivity Analysis (SPSA)"
        _style_title(ws3["A1"], size=16)
        ws3.merge_cells("A1:G1")
        ws3["A2"] = "Napolitano & Fabbri (1996). A negative difference means the effective weight is below the theoretical one at this site, not that the parameter is unimportant."
        ws3["A2"].font = Font(italic=True, size=9, color="666666")
        ws3.merge_cells("A2:G2")
        for col, h in enumerate(["Parameter", "Wi", "Ri", "Wi x Ri", "Theoretical %", "Effective %", "Difference %"], 1):
            _style_header(ws3.cell(row=4, column=col, value=h), bg_color=COLOR_ACCENT, font_color=COLOR_PRIMARY, size=11)
        row = 5
        for p in spsa_result["parameters"]:
            for c, v in enumerate([p["param"], p["weight"], p["rating"], p["weighted_score"], p["theoretical_pct"],
                                   p["effective_pct"], p["difference"]], 1):
                cell = ws3.cell(row=row, column=c, value=v)
                _style_value(cell)
                if c == 7:
                    cell.number_format = "+0.00;-0.00;0.00"
                    cell.font = Font(bold=True, color="2E7D32" if p["difference"] > 0 else "C62828" if p["difference"] < 0 else "000000")
            row += 1
        row = _kv(ws3, row + 1, "Most Influential", spsa_result.get("most_influential", "-"))
        _kv(ws3, row, "Base Index (DI)", spsa_result.get("base_index", "-"))
        for col, w in (("A", 18), ("B", 10), ("C", 10), ("D", 12), ("E", 16), ("F", 14), ("G", 14)):
            ws3.column_dimensions[col].width = w

    # ------------------------------------------------------------ 4. Monte Carlo
    if mc_result:
        ws4 = wb.create_sheet("Monte Carlo")
        ws4["A1"] = "Monte Carlo Uncertainty Analysis"
        _style_title(ws4["A1"], size=16)
        ws4.merge_cells("A1:C1")
        ws4["A2"] = (f"Iterations: {mc_result.get('n_iter', 'N/A')} | Variation: {mc_result.get('variation_pct', 'N/A')}% | "
                     "ASSUMED input uncertainty (user-chosen distributions), not field uncertainty.")
        ws4["A2"].font = Font(italic=True, size=9, color="666666")
        ws4.merge_cells("A2:C2")
        row = _section(ws4, 4, "DRASTIC", cols=3)
        ci = mc_result.get("ci_90", ("-", "-"))
        for label, val in (("Mean", mc_result.get("mean", "-")), ("Std Dev", mc_result.get("std", "-")),
                           ("Minimum", mc_result.get("min", "-")), ("Maximum", mc_result.get("max", "-")),
                           ("CI 90%", f"{ci[0]} - {ci[1]}"), ("P(index >= 140) %", mc_result.get("prob_over_140", "-"))):
            row = _kv(ws4, row, label, val)
        if "tox_mean" in mc_result:
            row = _section(ws4, row + 1, "DRASTIC-Tox (experimental)", cols=3)
            cit = mc_result.get("tox_ci_90", ("-", "-"))
            for label, val in (("Mean", mc_result.get("tox_mean", "-")), ("Std Dev", mc_result.get("tox_std", "-")),
                               ("Minimum", mc_result.get("tox_min", "-")), ("Maximum", mc_result.get("tox_max", "-")),
                               ("CI 90%", f"{cit[0]} - {cit[1]}"), ("P(index >= 140) %", mc_result.get("tox_prob_over_140", "-")),
                               ("P(index >= 180) %", mc_result.get("tox_prob_over_180", "-"))):
                row = _kv(ws4, row, label, val)
        ws4.column_dimensions["A"].width = 24
        ws4.column_dimensions["B"].width = 28
        ws4.column_dimensions["C"].width = 15

    # ------------------------------------------------------------ 5. Metadata
    ws5 = wb.create_sheet("Metadata")
    ws5["A1"] = "Report Metadata"
    _style_title(ws5["A1"], size=16)
    ws5.merge_cells("A1:B1")
    meta = [
        ("Application", "DRASTIC-Tox (Sudan Mining System)"), ("Version", "v61"),
        ("Report Date", now.strftime("%Y-%m-%d %H:%M:%S")),
        ("Institution", "University of Khartoum - Faculty of Engineering"), ("Status", "Pilot version"), ("", ""),
        ("References", "Aller et al. (1987) - DRASTIC, EPA/600/2-87/035"),
        ("", "Napolitano & Fabbri (1996) - SPSA"),
        ("", "Council Directive 98/83/EC - drinking-water limits used (CN 0.05, Hg 0.001 mg/L)"),
        ("", "WHO Guidelines for Drinking-water Quality, 4th ed. + addenda (Hg 0.006 mg/L; cyanide GV withdrawn)"),
        ("", "Konate et al. (2025) - modified DRASTIC for artisanal gold mining (context; DRASTIC-Tox coefficients are NOT taken from it)"),
        ("", ""),
        ("Disclaimer", "Screening tool only. DRASTIC-Tox is an experimental, externally unvalidated formulation. Does not replace laboratory analysis or professional judgement."),
    ]
    row = 3
    for label, val in meta:
        if label:
            ws5.cell(row=row, column=1, value=label)
            _style_label(ws5.cell(row=row, column=1))
        ws5.cell(row=row, column=2, value=_safe(val))
        _style_value(ws5.cell(row=row, column=2))
        row += 1
    ws5.column_dimensions["A"].width = 20
    ws5.column_dimensions["B"].width = 90

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
